"""Extrae cobertura documental; nunca configura tareas ni modifica el Excel.

Ejecutar con Python y openpyxl del runtime de documentos de Codex.
"""
import argparse
import csv
import hashlib
import json
from pathlib import Path

import openpyxl


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    source_bytes = args.source.read_bytes()
    workbook = openpyxl.load_workbook(args.source, data_only=False)
    sheet = workbook["Tareas"]
    assert sheet.cell(5, 2).value == "Nombre", "Encabezados inesperados"
    records = []
    for row in sheet.iter_rows(min_row=6):
        code = row[0].value
        if not code:
            continue
        assert isinstance(code, str) and code.startswith("TAR-"), "Código inesperado"
        values = [cell.value for cell in row]
        name = values[1]
        assert isinstance(name, str) and name.strip(), "Nombre faltante"
        filled = sum(value is not None for value in values[2:])
        comments = [(cell.coordinate, cell.comment.text) for cell in row if cell.comment]
        records.append({
            "CodigoOrigen": code,
            "Nombre": name,
            "Hoja": sheet.title,
            "FilaExcel": row[0].row,
            "Clasificacion": "PENDIENTE_DE_DEFINICION",
            "CamposOperativosLlenos": filled,
            "DudasOrigen": values[25] or "",
            "ComentariosOrigen": json.dumps(comments, ensure_ascii=False) if comments else "",
            "Motivo": "Falta confirmar si hay cambio de responsable, control intermedio, resultado conservado o plazo separado.",
        })
    codes = [record["CodigoOrigen"] for record in records]
    assert len(codes) == len(set(codes)) == 202, "Se esperan 202 códigos únicos"
    assert set(codes) == {f"TAR-{i:04}" for i in range(1, 203)}, "Cobertura incompleta"
    args.destination.mkdir(parents=True, exist_ok=True)
    csv_path = args.destination / "clasificacion-catalogo.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)
    evidence = {
        "source": str(args.source),
        "sha256": hashlib.sha256(source_bytes).hexdigest(),
        "tasks": len(records),
        "pending": len(records),
        "operational_cells_filled": sum(r["CamposOperativosLlenos"] for r in records),
        "cell_comments": sum(cell.comment is not None for s in workbook for row in s for cell in row),
        "classification": "Pendiente; no se deducen reglas a partir del nombre.",
    }
    (args.destination / "evidencia-catalogo.json").write_text(
        json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    assert args.source.read_bytes() == source_bytes, "El libro de origen cambió"
    print(json.dumps(evidence, ensure_ascii=False))


if __name__ == "__main__":
    main()
